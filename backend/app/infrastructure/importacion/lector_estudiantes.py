import csv
import io
import unicodedata
from typing import Dict, List

from app.application.dto.estudiantes import (
    COLUMNAS_IMPORTACION, COLUMNAS_OBLIGATORIAS, MAX_BYTES_IMPORTACION as MAX_BYTES, MAX_FILAS_IMPORTACION as MAX_FILAS,
)
from app.application.ports.lector_estudiantes import LectorEstudiantesPuerto
from app.domain.exceptions import ReglaDeNegocioViolada

_ALIAS = {
    "no_control": "numero_control", "num_control": "numero_control", "matricula": "numero_control",
    "apellido_paterno": "primer_apellido", "apellido_materno": "segundo_apellido",
    "nombre": "nombres", "nombre_s": "nombres", "telefono_alumno": "telefono",
    "correo_electronico": "correo", "email": "correo",
    "padre_tutor": "tutor_nombre", "tutor": "tutor_nombre", "telefono_tutor": "tutor_telefono",
    "correo_tutor": "tutor_correo", "parentesco": "tutor_parentesco",
}


def _normalizar_encabezado(texto: str) -> str:
    sin_acentos = "".join(c for c in unicodedata.normalize("NFKD", str(texto)) if not unicodedata.combining(c))
    limpio = "".join(c if c.isalnum() else "_" for c in sin_acentos.strip().lower())
    limpio = "_".join(p for p in limpio.split("_") if p)
    return _ALIAS.get(limpio, limpio)


def _texto_celda(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))  # Excel guarda 4491234567 como 4491234567.0
    return str(valor).strip()


class LectorEstudiantes(LectorEstudiantesPuerto):
    def leer(self, contenido: bytes, nombre_archivo: str) -> List[Dict[str, str]]:
        if len(contenido) > MAX_BYTES:
            raise ReglaDeNegocioViolada("El archivo excede el tamaño máximo de 2 MB.")
        nombre = (nombre_archivo or "").lower()
        if nombre.endswith(".csv"):
            filas = self._filas_csv(contenido)
        elif nombre.endswith(".xlsx"):
            filas = self._filas_xlsx(contenido)
        else:
            raise ReglaDeNegocioViolada("Formato no soportado. Usa un archivo .csv o .xlsx.")
        return self._a_diccionarios(filas)

    # --- formatos ---
    def _filas_csv(self, contenido: bytes):
        try:
            texto = contenido.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = contenido.decode("cp1252")  # CSV guardado por Excel en Windows
        primera = texto.splitlines()[0] if texto.strip() else ""
        delimitador = ";" if primera.count(";") > primera.count(",") else ","
        return [[c.strip() for c in fila] for fila in csv.reader(io.StringIO(texto), delimiter=delimitador)]

    def _filas_xlsx(self, contenido: bytes):
        from openpyxl import load_workbook  # import diferido: solo se necesita para .xlsx

        try:
            libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
            hoja = libro.active
            return [[_texto_celda(c) for c in fila] for fila in hoja.iter_rows(values_only=True)]
        except Exception as exc:  # archivo corrupto o no es un xlsx válido
            raise ReglaDeNegocioViolada("No se pudo leer el archivo Excel.") from exc

    # --- estructura ---
    def _a_diccionarios(self, filas) -> List[Dict[str, str]]:
        filas = [f for f in filas if any(str(c).strip() for c in f)]
        if not filas:
            raise ReglaDeNegocioViolada("El archivo está vacío.")
        encabezados = [_normalizar_encabezado(h) for h in filas[0]]
        faltantes = [c for c in COLUMNAS_OBLIGATORIAS if c not in encabezados]
        if faltantes:
            raise ReglaDeNegocioViolada("Faltan columnas obligatorias: " + ", ".join(faltantes) + ".")
        datos = filas[1:]
        if len(datos) > MAX_FILAS:
            raise ReglaDeNegocioViolada(f"El archivo tiene más de {MAX_FILAS} filas. Divídelo en partes.")
        resultado = []
        for fila in datos:
            registro = {c: "" for c in COLUMNAS_IMPORTACION}
            for nombre, valor in zip(encabezados, fila):
                if nombre in registro:
                    registro[nombre] = _texto_celda(valor)
            resultado.append(registro)
        return resultado

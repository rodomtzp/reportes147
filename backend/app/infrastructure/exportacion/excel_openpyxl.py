"""Generador del .xlsx (openpyxl): hoja «Resumen» y hoja «Reportes»."""
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.application.dto.exportaciones import LibroReportes
from app.application.ports.exportador import ExportadorExcelPuerto

_ENCABEZADO = PatternFill("solid", fgColor="1F3A5F")
_PREFIJOS_FORMULA = ("=", "+", "-", "@", "\t", "\r")


def texto_seguro(valor: Any) -> Any:
    """Evita inyección de fórmulas: un texto que Excel interpretaría como fórmula se guarda como texto."""
    if isinstance(valor, str) and valor.startswith(_PREFIJOS_FORMULA):
        return "'" + valor
    return valor


class ExportadorExcelOpenpyxl(ExportadorExcelPuerto):
    def generar(self, libro: LibroReportes) -> bytes:
        wb = Workbook()
        self._resumen(wb.active, libro)
        self._reportes(wb.create_sheet("Reportes"), libro)
        salida = BytesIO()
        wb.save(salida)
        return salida.getvalue()

    # ---- hojas ----
    @staticmethod
    def _resumen(ws, libro: LibroReportes) -> None:
        ws.title = "Resumen"
        ws["A1"], ws["A1"].font = libro.plantel, Font(bold=True, size=14)
        ws["A2"], ws["A2"].font = libro.titulo, Font(bold=True, size=12)
        ws["A3"] = texto_seguro(libro.subtitulo)
        ws["A4"] = f"Generado por {texto_seguro(libro.generado_por)} el {libro.generado_en:%d/%m/%Y %H:%M}"
        ws["A5"] = "Incluye datos de contacto" if libro.incluye_contacto else "Sin datos de contacto"
        vigentes = sum(1 for f in libro.filas if not f.cancelado)
        fila = 7

        def bloque(titulo: str, pares) -> None:
            nonlocal fila
            ws.cell(fila, 1, titulo).font = Font(bold=True, color="FFFFFF")
            ws.cell(fila, 1).fill = ws.cell(fila, 2).fill = _ENCABEZADO
            fila += 1
            for etiqueta, valor in pares:
                ws.cell(fila, 1, texto_seguro(etiqueta))
                ws.cell(fila, 2, valor)
                fila += 1
            fila += 1

        bloque("Totales", [
            ("Reportes (no cancelados)", vigentes),
            ("Reportes cancelados", len(libro.filas) - vigentes),
            ("Sin llamada efectiva (pendiente o no contestó)", libro.sin_llamada_efectiva),
        ])
        bloque("Por gravedad", libro.por_gravedad)
        bloque("Por estado", libro.por_estado)
        bloque("Sanciones aplicadas", libro.por_sancion or [("Ninguna", 0)])
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 14

    @staticmethod
    def _reportes(ws, libro: LibroReportes) -> None:
        columnas = [
            ("Folio", 15, "folio"), ("Fecha del incidente", 18, "fecha_incidente"), ("No. de control", 17, "numero_control"),
            ("Estudiante", 32, "estudiante"), ("Grupo", 8, "grupo"), ("Carrera", 26, "carrera"),
            ("Incidencia", 26, "tipo_incidencia"), ("Descripción", 50, "descripcion"), ("Gravedad", 11, "gravedad"),
            ("Estado", 26, "estado"), ("Cancelado", 11, "cancelado"), ("Motivo de cancelación", 30, "motivo_cancelacion"),
            ("Cuenta para acumulación", 14, "cuenta_acumulacion"), ("Reportado por", 26, "creado_por"),
            ("Llamada", 13, "llamada"), ("Sanción", 34, "sancion"),
        ]
        if libro.incluye_contacto:
            columnas += [("Padre o tutor", 28, "tutor_legal"), ("Teléfono", 14, "telefono_tutor"),
                         ("Correo del tutor", 28, "correo_tutor"), ("Domicilio", 40, "domicilio")]
        for i, (titulo, ancho, _) in enumerate(columnas, 1):
            c = ws.cell(1, i, titulo)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), _ENCABEZADO
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = ancho
        for n, f in enumerate(libro.filas, 2):
            for i, (_, _, campo) in enumerate(columnas, 1):
                valor = getattr(f, campo)
                if campo == "fecha_incidente":
                    valor = valor.replace(tzinfo=None)
                    celda = ws.cell(n, i, valor)
                    celda.number_format = "dd/mm/yyyy hh:mm"
                    continue
                if isinstance(valor, bool):
                    valor = "Sí" if valor else "No"
                celda = ws.cell(n, i, texto_seguro(valor))
                if campo in ("descripcion", "motivo_cancelacion", "domicilio"):
                    celda.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{max(len(libro.filas) + 1, 2)}"

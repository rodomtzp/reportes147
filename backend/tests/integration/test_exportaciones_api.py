from io import BytesIO

from openpyxl import load_workbook

from tests.integration.test_reportes_api import _cuerpo, ctx  # noqa: F401  (ctx es una fixture)

BASE = "/api/v1"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _crear(client, h, ctx, letras="a"):
    assert client.post(f"{BASE}/reportes", headers=h, json=_cuerpo(ctx, letras)).status_code == 201


def _hoja(respuesta, nombre="Reportes"):
    return load_workbook(BytesIO(respuesta.data))[nombre]


def test_exporta_por_estudiante_grupo_y_carrera(client, ctx, headers_de):
    docente, orient = headers_de("DOCENTE"), headers_de("ORIENTACION")
    _crear(client, docente, ctx, "ab")
    _crear(client, docente, ctx, "d")

    r = client.get(f"{BASE}/exportaciones/estudiantes/{ctx['est']['a']}", headers=orient)
    assert r.status_code == 200 and r.mimetype == XLSX
    assert "attachment" in r.headers["Content-Disposition"] and ".xlsx" in r.headers["Content-Disposition"]
    assert r.headers["Cache-Control"] == "private, no-store"
    assert _hoja(r).max_row == 2

    por_grupo = client.get(f"{BASE}/exportaciones/grupos/{ctx['grupos']['1G']['id']}", headers=orient)
    assert _hoja(por_grupo).max_row == 3
    por_carrera = client.get(f"{BASE}/exportaciones/carreras/{ctx['grupos']['1G']['carrera_id']}", headers=orient)
    assert _hoja(por_carrera).max_row == 4
    assert _hoja(por_carrera, "Resumen")["A2"].value == "Reportes disciplinarios"


def test_permisos_y_errores_de_exportacion(client, ctx, headers_de):
    docente, prefecto, admin = headers_de("DOCENTE"), headers_de("PREFECTO"), ctx["admin"]
    _crear(client, docente, ctx, "a")
    url = f"{BASE}/exportaciones/estudiantes/{ctx['est']['a']}"
    assert client.get(url).status_code == 401
    assert client.get(url, headers=prefecto).status_code == 403
    assert client.get(url, headers=admin).status_code == 403          # el Administrador no consulta reportes
    assert client.get(url, headers=docente).status_code == 200        # el docente, los que creó
    orient = headers_de("ORIENTACION")
    for ruta in ("estudiantes", "grupos", "carreras"):
        assert client.get(f"{BASE}/exportaciones/{ruta}/99999", headers=orient).status_code == 404
    assert client.get(f"{url}?desde=ayer", headers=orient).status_code == 422


def test_contacto_opcional_y_con_permiso(client, ctx, headers_de):
    docente, orient = headers_de("DOCENTE"), headers_de("ORIENTACION")
    _crear(client, docente, ctx, "a")
    url = f"{BASE}/exportaciones/estudiantes/{ctx['est']['a']}"
    assert "Teléfono" not in [c.value for c in _hoja(client.get(url, headers=orient))[1]]
    con = _hoja(client.get(f"{url}?incluir_contacto=true", headers=orient))
    assert "Teléfono" in [c.value for c in con[1]]
    assert client.get(f"{url}?incluir_contacto=true", headers=docente).status_code == 403


def test_bitacora_solo_para_el_administrador(client, ctx, headers_de):
    docente, orient = headers_de("DOCENTE"), headers_de("ORIENTACION")
    _crear(client, docente, ctx, "a")
    client.get(f"{BASE}/exportaciones/grupos/{ctx['grupos']['1G']['id']}?incluir_contacto=true", headers=orient)
    client.get(f"{BASE}/estudiantes/{ctx['est']['a']}", headers=orient)           # ver la ficha con contacto

    for h in (docente, orient, headers_de("SUBDIRECCION"), headers_de("DIRECCION")):
        assert client.get(f"{BASE}/bitacora", headers=h).status_code == 403
    registros = client.get(f"{BASE}/bitacora", headers=ctx["admin"]).get_json()
    acciones = {r["accion"] for r in registros}
    assert acciones == {"EXPORTACION", "CONSULTA_CONTACTO"}
    exp = next(r for r in registros if r["accion"] == "EXPORTACION")
    assert exp["filas"] == 1 and exp["con_contacto"] is True and exp["usuario"]
    solo = client.get(f"{BASE}/bitacora?accion=EXPORTACION", headers=ctx["admin"]).get_json()
    assert [r["accion"] for r in solo] == ["EXPORTACION"]
    assert client.get(f"{BASE}/bitacora?accion=XX", headers=ctx["admin"]).status_code == 422

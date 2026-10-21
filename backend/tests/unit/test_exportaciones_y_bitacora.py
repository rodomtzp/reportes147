from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.application.ports.exportador import ExportadorExcelPuerto
from app.application.use_cases.bitacora.consultar import ConsultarBitacora
from app.application.use_cases.estudiantes.consultas import ObtenerEstudiante
from app.application.use_cases.exportaciones.exportar_reportes import ExportarReportes
from app.application.use_cases.llamadas.seguimiento import ObtenerSeguimientoLlamadas
from app.application.use_cases.reportes.crear_reportes import CrearReportes
from app.domain.exceptions import NoEncontrado, PermisoDenegado, ReglaDeNegocioViolada
from app.domain.policies.politica_permisos import Permiso, tiene_permiso
from app.domain.value_objects.accion_bitacora import AccionBitacora
from app.domain.value_objects.gravedad import Gravedad as G
from app.domain.value_objects.rol import Rol
from app.infrastructure.exportacion.excel_openpyxl import ExportadorExcelOpenpyxl, texto_seguro

TEXTO = "Se peleó con un compañero durante el receso."


class ExportadorEspia(ExportadorExcelPuerto):
    def __init__(self):
        self.libro = None

    def generar(self, libro):
        self.libro = libro
        return b"xlsx"


@pytest.fixture
def espia():
    return ExportadorEspia()


@pytest.fixture
def datos(uow, reloj, notificador, mundo):
    """Reportes: a (docente), a (docente2), b (docente), d (docente, grupo 1H)."""
    def crear(letra, autor, gravedad=G.LEVE):
        reloj.avanzar(minutes=1)
        return CrearReportes(uow, reloj, notificador).ejecutar(
            mundo["u"][autor], [mundo["est"][letra].id], mundo["pelea"].id, TEXTO, gravedad)[0].reporte_id

    return {"a1": crear("a", "docente"), "a2": crear("a", "docente2", G.GRAVE), "b": crear("b", "docente"),
            "d": crear("d", "docente")}


def _exportar(uow, reloj, espia, actor, alcance, objetivo, **kw):
    return ExportarReportes(uow, reloj, espia, "CBTis 147").ejecutar(actor, alcance, objetivo, **kw)


def test_permiso_de_exportar_incluye_al_docente_y_la_bitacora_es_del_administrador():
    assert tiene_permiso([Rol.DOCENTE], Permiso.EXPORTAR_EXCEL)
    assert not tiene_permiso([Rol.PREFECTO], Permiso.EXPORTAR_EXCEL)
    assert tiene_permiso([Rol.ADMINISTRADOR], Permiso.CONSULTAR_BITACORA)
    assert not tiene_permiso([Rol.SUBDIRECCION, Rol.DIRECCION, Rol.ORIENTACION], Permiso.CONSULTAR_BITACORA)


def test_directivo_exporta_por_estudiante_grupo_y_carrera(uow, reloj, espia, mundo, datos):
    o = mundo["u"]["orientacion"]
    a = _exportar(uow, reloj, espia, o, "estudiante", mundo["est"]["a"].id)
    assert a.filas == 2 and a.nombre.startswith("reportes_estudiante_") and a.nombre.endswith(".xlsx")
    assert _exportar(uow, reloj, espia, o, "grupo", mundo["grupos"]["1G"].id).filas == 3
    assert _exportar(uow, reloj, espia, o, "grupo", mundo["grupos"]["1H"].id).filas == 1
    assert _exportar(uow, reloj, espia, o, "carrera", mundo["carrera"].id).filas == 4


def test_resumen_cuenta_por_gravedad_y_llamadas(uow, reloj, espia, mundo, datos):
    _exportar(uow, reloj, espia, mundo["u"]["orientacion"], "carrera", mundo["carrera"].id)
    libro = espia.libro
    assert dict(libro.por_gravedad) == {"Leve": 3, "Media": 0, "Grave": 1, "Muy grave": 0}
    assert libro.sin_llamada_efectiva == 4 and libro.plantel == "CBTis 147"


def test_docente_solo_exporta_lo_que_creo_y_el_tutor_su_grupo(uow, reloj, espia, mundo, datos):
    docente, tutor = mundo["u"]["docente"], mundo["u"]["tutor"]
    assert _exportar(uow, reloj, espia, docente, "estudiante", mundo["est"]["a"].id).filas == 1  # solo el suyo
    assert _exportar(uow, reloj, espia, docente, "carrera", mundo["carrera"].id).filas == 3
    assert _exportar(uow, reloj, espia, tutor, "grupo", mundo["grupos"]["1G"].id).filas == 3  # todo el grupo
    assert _exportar(uow, reloj, espia, tutor, "grupo", mundo["grupos"]["1H"].id).filas == 0  # no es suyo
    with pytest.raises(PermisoDenegado):
        _exportar(uow, reloj, espia, mundo["u"]["prefecto"], "grupo", mundo["grupos"]["1G"].id)


def test_contacto_solo_si_se_pide_y_se_tiene_permiso(uow, reloj, espia, mundo, datos):
    o, tutor, docente = mundo["u"]["orientacion"], mundo["u"]["tutor"], mundo["u"]["docente2"]
    _exportar(uow, reloj, espia, o, "estudiante", mundo["est"]["a"].id)
    assert espia.libro.incluye_contacto is False and espia.libro.filas[0].telefono_tutor is None

    _exportar(uow, reloj, espia, o, "estudiante", mundo["est"]["a"].id, incluir_contacto=True)
    f = espia.libro.filas[0]
    assert espia.libro.incluye_contacto and f.telefono_tutor == "4497654321" and f.tutor_legal == "Laura Gómez (Madre)"

    _exportar(uow, reloj, espia, tutor, "grupo", mundo["grupos"]["1G"].id, incluir_contacto=True)
    assert all(f.telefono_tutor for f in espia.libro.filas)          # su grupo: sí

    with pytest.raises(PermisoDenegado):                                 # docente sin grupos: no
        _exportar(uow, reloj, espia, docente, "estudiante", mundo["est"]["a"].id, incluir_contacto=True)


def test_tutor_no_ve_contacto_de_estudiantes_fuera_de_su_grupo(uow, reloj, espia, mundo, datos):
    """El tutor de 1G es además autor de un reporte de 1H: ese estudiante sale sin contacto."""
    tutor = mundo["u"]["tutor"]
    CrearReportes(uow, reloj, espia and __import__("tests.unit.fakes", fromlist=["NotificadorMem"]).NotificadorMem()).ejecutar(
        tutor, [mundo["est"]["d"].id], mundo["pelea"].id, TEXTO, G.LEVE)
    _exportar(uow, reloj, espia, tutor, "carrera", mundo["carrera"].id, incluir_contacto=True)
    por_grupo = {f.grupo: f.telefono_tutor for f in espia.libro.filas}
    assert por_grupo["1G"] and por_grupo["1H"] is None


def test_la_exportacion_queda_en_bitacora_y_la_consulta_solo_el_administrador_via_api(uow, reloj, espia, mundo, datos):
    o = mundo["u"]["orientacion"]
    _exportar(uow, reloj, espia, o, "grupo", mundo["grupos"]["1G"].id, incluir_contacto=True)
    (r,) = ConsultarBitacora(uow).ejecutar(accion=AccionBitacora.EXPORTACION)
    assert (r.registro.usuario_id, r.registro.filas, r.registro.con_contacto) == (o.id, 3, True)
    assert r.usuario_nombre == "Orientación" and "1G" in r.registro.descripcion
    assert uow.commits >= 1


def test_consultar_contacto_queda_en_bitacora(uow, reloj, mundo, datos):
    o = mundo["u"]["orientacion"]
    ObtenerEstudiante(uow, reloj).ejecutar(o, mundo["est"]["a"].id)
    ObtenerSeguimientoLlamadas(uow, reloj).ejecutar(o.id, datos["a1"])
    registros = ConsultarBitacora(uow).ejecutar(accion=AccionBitacora.CONSULTA_CONTACTO, estudiante_id=mundo["est"]["a"].id)
    assert len(registros) == 2 and all(r.estudiante == mundo["est"]["a"].nombre_completo for r in registros)
    assert ConsultarBitacora(uow).ejecutar(accion=AccionBitacora.EXPORTACION) == []
    assert ConsultarBitacora(uow).ejecutar(usuario_id=o.id + 99) == []


def test_valida_alcance_y_objetivo(uow, reloj, espia, mundo, datos):
    o = mundo["u"]["orientacion"]
    with pytest.raises(ReglaDeNegocioViolada):
        _exportar(uow, reloj, espia, o, "plantel", 1)
    for alcance in ("estudiante", "grupo", "carrera"):
        with pytest.raises(NoEncontrado):
            _exportar(uow, reloj, espia, o, alcance, 9999)


def test_texto_seguro_neutraliza_formulas():
    assert texto_seguro("=HYPERLINK(\"http://x\")") == "'=HYPERLINK(\"http://x\")"
    assert texto_seguro("+52") == "'+52" and texto_seguro("@cmd") == "'@cmd"
    assert texto_seguro("Normal") == "Normal" and texto_seguro(5) == 5


def test_excel_real_tiene_hojas_resumen_y_reportes(uow, reloj, mundo, datos):
    archivo = ExportarReportes(uow, reloj, ExportadorExcelOpenpyxl(), "CBTis 147").ejecutar(
        mundo["u"]["orientacion"], "grupo", mundo["grupos"]["1G"].id, incluir_contacto=True)
    wb = load_workbook(BytesIO(archivo.contenido))
    assert wb.sheetnames == ["Resumen", "Reportes"]
    ws = wb["Reportes"]
    encabezados = [c.value for c in ws[1]]
    assert encabezados[0] == "Folio" and "Teléfono" in encabezados
    assert ws.max_row == 4                                    # encabezado + 3 reportes
    assert str(ws.cell(2, 1).value).startswith("RPT-2026-")
    resumen = {r[0].value: r[1].value for r in wb["Resumen"].iter_rows(min_row=7) if r[0].value}
    assert resumen["Reportes (no cancelados)"] == 3 and resumen["Grave"] == 1
